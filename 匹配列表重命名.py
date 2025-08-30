import os
import shutil
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook
import threading
from datetime import datetime
import difflib
from collections import defaultdict

class PaperRenamerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("论文文件重命名工具")
        self.root.geometry("800x600")
        
        # 变量初始化
        self.excel_path = tk.StringVar()
        self.folder_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.unmatched_files = []  # 存储未匹配的文件列表
        self.similarity_threshold = 0.6  # 默认相似度阈值
        
        self.setup_ui()
        
    def setup_ui(self):
        # 创建主框架
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        # Excel文件选择部分
        ttk.Label(main_frame, text="Excel文件路径:").grid(row=0, column=0, sticky=tk.W, pady=5)
        ttk.Entry(main_frame, textvariable=self.excel_path, width=60).grid(row=0, column=1, padx=5)
        ttk.Button(main_frame, text="浏览", command=self.select_excel).grid(row=0, column=2, padx=5)
        
        # 源文件夹选择部分
        ttk.Label(main_frame, text="论文文件夹路径:").grid(row=1, column=0, sticky=tk.W, pady=5)
        ttk.Entry(main_frame, textvariable=self.folder_path, width=60).grid(row=1, column=1, padx=5)
        ttk.Button(main_frame, text="浏览", command=self.select_folder).grid(row=1, column=2, padx=5)
        
        # 输出文件夹选择部分
        ttk.Label(main_frame, text="输出文件夹路径:").grid(row=2, column=0, sticky=tk.W, pady=5)
        ttk.Entry(main_frame, textvariable=self.output_path, width=60).grid(row=2, column=1, padx=5)
        ttk.Button(main_frame, text="浏览", command=self.select_output_folder).grid(row=2, column=2, padx=5)
        
        # 表格列设置
        ttk.Label(main_frame, text="序号列名:").grid(row=3, column=0, sticky=tk.W, pady=5)
        self.id_column = ttk.Entry(main_frame, width=20)
        self.id_column.grid(row=3, column=1, sticky=tk.W, padx=5)
        self.id_column.insert(0, "序号")  # 默认值
        
        ttk.Label(main_frame, text="论文题目列名:").grid(row=4, column=0, sticky=tk.W, pady=5)
        self.title_column = ttk.Entry(main_frame, width=20)
        self.title_column.grid(row=4, column=1, sticky=tk.W, padx=5)
        self.title_column.insert(0, "论文题目")  # 默认值
        
        # 相似度阈值设置
        ttk.Label(main_frame, text="相似度阈值 (0.1-1.0):").grid(row=5, column=0, sticky=tk.W, pady=5)
        self.similarity_entry = ttk.Entry(main_frame, width=10)
        self.similarity_entry.grid(row=5, column=1, sticky=tk.W, padx=5)
        self.similarity_entry.insert(0, "0.6")  # 默认值
        
        # 执行按钮
        ttk.Button(main_frame, text="开始处理", command=self.start_processing).grid(row=6, column=1, pady=20)
        
        # 进度条
        self.progress = ttk.Progressbar(main_frame, mode='indeterminate')
        self.progress.grid(row=7, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=5)
        
        # 日志文本框
        ttk.Label(main_frame, text="处理日志:").grid(row=8, column=0, sticky=tk.W, pady=5)
        
        # 添加滚动条
        log_frame = ttk.Frame(main_frame)
        log_frame.grid(row=9, column=0, columnspan=3, sticky=(tk.W, tk.E, tk.N, tk.S), pady=5)
        
        self.log_text = tk.Text(log_frame, height=20, width=90)
        scrollbar = ttk.Scrollbar(log_frame, orient="vertical", command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)
        
        self.log_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
        
        # 配置网格权重
        main_frame.columnconfigure(1, weight=1)
        main_frame.rowconfigure(9, weight=1)
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        
    def select_excel(self):
        file_path = filedialog.askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        if file_path:
            self.excel_path.set(file_path)
            
    def select_folder(self):
        folder_path = filedialog.askdirectory(title="选择论文文件夹")
        if folder_path:
            self.folder_path.set(folder_path)
            
    def select_output_folder(self):
        folder_path = filedialog.askdirectory(title="选择输出文件夹")
        if folder_path:
            self.output_path.set(folder_path)
            
    def log_message(self, message):
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.root.update_idletasks()
        
    def start_processing(self):
        if not self.excel_path.get() or not self.folder_path.get() or not self.output_path.get():
            messagebox.showerror("错误", "请先选择Excel文件、论文文件夹和输出文件夹")
            return
            
        # 获取相似度阈值
        try:
            threshold = float(self.similarity_entry.get().strip())
            if threshold < 0.1 or threshold > 1.0:
                messagebox.showerror("错误", "相似度阈值必须在0.1到1.0之间")
                return
            self.similarity_threshold = threshold
        except ValueError:
            messagebox.showerror("错误", "请输入有效的相似度阈值（0.1-1.0）")
            return
            
        # 在后台线程中执行处理，避免界面冻结
        thread = threading.Thread(target=self.process_files)
        thread.daemon = True
        thread.start()
        
    def find_best_match(self, filename, paper_titles):
        """使用模糊匹配找到最佳匹配的论文标题"""
        best_match = None
        best_score = 0
        
        for title in paper_titles:
            # 计算相似度
            score = difflib.SequenceMatcher(None, filename.lower(), title.lower()).ratio()
            
            # 如果相似度高于阈值且是当前最佳匹配
            if score > best_score and score >= self.similarity_threshold:
                best_score = score
                best_match = title
        
        return best_match, best_score
    
    def show_unmatched_files_popup(self, unmatched_count, unmatched_files):
        """显示未匹配文件的弹窗提示"""
        if unmatched_count == 0:
            messagebox.showinfo("处理完成", f"所有文件都已成功匹配并重命名！\n\n输出文件夹: {self.output_path.get()}")
            return
            
        # 创建弹窗
        popup = tk.Toplevel(self.root)
        popup.title("未匹配文件提示")
        popup.geometry("600x400")
        popup.transient(self.root)
        popup.grab_set()
        
        # 设置弹窗内容
        label = ttk.Label(popup, text=f"有 {unmatched_count} 个文件未能匹配:", font=("Arial", 12))
        label.pack(pady=10)
        
        # 创建文本框显示未匹配文件列表
        text_frame = ttk.Frame(popup)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        text_widget = tk.Text(text_frame, wrap=tk.WORD)
        scrollbar = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        text_widget.configure(yscrollcommand=scrollbar.set)
        
        text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        for file in unmatched_files:
            text_widget.insert(tk.END, f"• {file}\n")
        
        text_widget.config(state=tk.DISABLED)  # 设置为只读
        
        # 确定按钮
        button_frame = ttk.Frame(popup)
        button_frame.pack(pady=10)
        
        ttk.Button(button_frame, text="确定", command=popup.destroy).pack()
        
    def process_files(self):
        try:
            self.progress.start()
            self.log_message("开始处理...")
            self.log_message(f"使用相似度阈值: {self.similarity_threshold}")
            
            # 清空未匹配文件列表
            self.unmatched_files = []
            
            # 读取Excel文件
            excel_file = self.excel_path.get()
            id_col = self.id_column.get().strip()
            title_col = self.title_column.get().strip()
            
            try:
                # 尝试多种方式读取Excel
                try:
                    df = pd.read_excel(excel_file)
                except:
                    wb = load_workbook(excel_file)
                    sheet = wb.active
                    data = sheet.values
                    cols = next(data)
                    df = pd.DataFrame(data, columns=cols)
                
                # 检查列是否存在
                if id_col not in df.columns or title_col not in df.columns:
                    self.log_message(f"错误: Excel文件中未找到列 '{id_col}' 或 '{title_col}'")
                    return
                    
                # 创建论文标题到序号的映射
                paper_map = {}
                paper_titles = []  # 存储所有论文标题用于模糊匹配
                for _, row in df.iterrows():
                    paper_id = str(row[id_col]).strip()
                    paper_title = str(row[title_col]).strip()
                    if paper_id and paper_title:
                        paper_map[paper_title] = paper_id
                        paper_titles.append(paper_title)
                
                self.log_message(f"成功读取 {len(paper_map)} 条论文信息")
                
            except Exception as e:
                self.log_message(f"读取Excel文件时出错: {str(e)}")
                return
                
            # 创建输出文件夹（如果不存在）
            output_folder = self.output_path.get()
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_subfolder = os.path.join(output_folder, f"重命名论文_{timestamp}")
            
            try:
                os.makedirs(output_subfolder, exist_ok=True)
                self.log_message(f"创建输出文件夹: {output_subfolder}")
            except Exception as e:
                self.log_message(f"创建输出文件夹时出错: {str(e)}")
                return
                
            # 处理文件夹中的文件
            folder_path = self.folder_path.get()
            files = os.listdir(folder_path)
            matched_count = 0
            fuzzy_matched_count = 0  # 模糊匹配计数
            
            for filename in files:
                file_path = os.path.join(folder_path, filename)
                if os.path.isfile(file_path):
                    # 获取文件名（不含扩展名）进行匹配
                    name_without_ext = os.path.splitext(filename)[0]
                    
                    # 先尝试精确匹配
                    if name_without_ext in paper_map:
                        # 精确匹配成功
                        paper_id = paper_map[name_without_ext]
                        file_ext = os.path.splitext(filename)[1]
                        new_filename = f"{paper_id}_{filename}"
                        new_file_path = os.path.join(output_subfolder, new_filename)
                        
                        try:
                            shutil.copy2(file_path, new_file_path)
                            self.log_message(f"精确匹配: {filename} -> {new_filename}")
                            matched_count += 1
                        except Exception as e:
                            self.log_message(f"复制文件 {filename} 时出错: {str(e)}")
                    else:
                        # 精确匹配失败，尝试模糊匹配
                        best_match, similarity_score = self.find_best_match(name_without_ext, paper_titles)
                        
                        if best_match:
                            # 模糊匹配成功
                            paper_id = paper_map[best_match]
                            file_ext = os.path.splitext(filename)[1]
                            new_filename = f"{paper_id}_{filename}"
                            new_file_path = os.path.join(output_subfolder, new_filename)
                            
                            try:
                                shutil.copy2(file_path, new_file_path)
                                self.log_message(f"模糊匹配 (相似度: {similarity_score:.2f}): {filename} -> {new_filename}")
                                self.log_message(f"  匹配论文: {best_match}")
                                matched_count += 1
                                fuzzy_matched_count += 1
                            except Exception as e:
                                self.log_message(f"复制文件 {filename} 时出错: {str(e)}")
                        else:
                            # 两种匹配都失败
                            self.unmatched_files.append(filename)
            
            # 输出统计信息
            self.log_message("\n处理完成!")
            self.log_message(f"成功匹配并重命名了 {matched_count} 个文件")
            if fuzzy_matched_count > 0:
                self.log_message(f"  - 其中 {fuzzy_matched_count} 个通过模糊匹配")
            self.log_message(f"输出文件夹: {output_subfolder}")
            self.log_message(f"未能匹配 {len(self.unmatched_files)} 个文件")
            
            if self.unmatched_files:
                self.log_message("\n未匹配的文件列表:")
                for file in self.unmatched_files:
                    self.log_message(f"  - {file}")
            
            # 检查是否有Excel中的论文没有对应文件
            matched_titles = set()
            for filename in files:
                name_without_ext = os.path.splitext(filename)[0]
                best_match, _ = self.find_best_match(name_without_ext, paper_titles)
                if best_match:
                    matched_titles.add(best_match)
            
            unmatched_papers = set(paper_titles) - matched_titles
            if unmatched_papers:
                self.log_message(f"\nExcel中有 {len(unmatched_papers)} 篇论文没有对应的文件:")
                for paper in unmatched_papers:
                    self.log_message(f"  - {paper}")
                    
            # 处理完成后显示弹窗提示
            self.root.after(0, lambda: self.show_unmatched_files_popup(
                len(self.unmatched_files), self.unmatched_files))
                    
        except Exception as e:
            self.log_message(f"处理过程中发生错误: {str(e)}")
        finally:
            self.progress.stop()
            self.log_message("处理结束")

if __name__ == "__main__":
    root = tk.Tk()
    app = PaperRenamerApp(root)
    root.mainloop()